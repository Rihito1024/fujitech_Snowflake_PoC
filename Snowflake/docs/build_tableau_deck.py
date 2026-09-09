#!/usr/bin/env python3
"""フジテック様向け Tableau Cloud 接続手順書を エクスチュア公式テンプレートで生成。

使い方（リポジトリ直下で実行）:
    pip install python-pptx
    python Snowflake/docs/build_tableau_deck.py

前提: リポジトリ直下に エクスチュア公式テンプレート `template.pptx` を置く
      （社内 Exportal から入手。リポジトリには含めていない）。
内容の元ネタは Snowflake/docs/tableau_cloud_connection_setup.md。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "template.pptx")
OUT = os.path.join(ROOT, "Snowflake/docs/Tableau_Cloud接続手順書_フジテック様.pptx")

GREEN = RGBColor(0x5E, 0x91, 0x3B)
GRAY_D = RGBColor(0x30, 0x30, 0x30)
GRAY_L = RGBColor(0xA5, 0xA5, 0xA5)
RED = RGBColor(0xC0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(SRC)

# ---- レイアウト取得（名前引き） -------------------------------------------------
L = {ly.name: ly for ly in prs.slide_masters[0].slide_layouts}
LY_COVER = L["表紙"]
LY_DIV = L["中表紙"]
LY_STD = L["標準ページ　下線あり"]
LY_OL = L["標準ページ　下線あり　OL対応1"]
LY_END = L["最終ページ"]

# ---- 既存サンプルスライドを全削除 ---------------------------------------------
xml_slides = prs.slides._sldIdLst
for sldId in list(xml_slides):
    rId = sldId.get(qn("r:id"))
    prs.part.drop_rel(rId)
    xml_slides.remove(sldId)


def add(layout):
    return prs.slides.add_slide(layout)


def set_ph(slide, idx, text):
    ph = slide.placeholders[idx]
    ph.text_frame.text = text
    return ph


def body_lines(slide, idx, items):
    """items: list of (level, text) or (level, text, bold)"""
    tf = slide.placeholders[idx].text_frame
    tf.clear()
    for i, it in enumerate(items):
        lvl, txt = it[0], it[1]
        bold = it[2] if len(it) > 2 else None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = txt
        p.level = lvl
        if bold is not None:
            for r in p.runs:
                r.font.bold = bold


def textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb


def screenshot_ph(slide, l, t, w, h, caption):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
    sh.line.color.rgb = GRAY_L; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "［スクリーンショット挿入予定］\n" + caption
    r.font.size = Pt(12); r.font.color.rgb = GRAY_L
    return sh


def make_table(slide, rows, l, t, w, h, col_w=None, header=True, first_col_accent=False):
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    if col_w:
        for ci, cw in enumerate(col_w):
            tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = str(val)
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11.5)
                    r.font.color.rgb = GRAY_D
            is_head = header and ri == 0
            is_fca = first_col_accent and ci == 0 and not is_head
            if is_head:
                cell.fill.solid(); cell.fill.fore_color.rgb = GREEN
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.bold = True; r.font.color.rgb = WHITE
            elif is_fca:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xEE, 0xF3, 0xE9)
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
    return tbl


# ============================================================================
# 1. 表紙
# ============================================================================
s = add(LY_COVER)
set_ph(s, 0, "Snowflake 接続手順書\nTableau Cloud からのデータ参照")
set_ph(s, 11, "フジテック株式会社　御中")
set_ph(s, 12, "2026年9月")

# ============================================================================
# 2. 目次
# ============================================================================
s = add(LY_OL)
set_ph(s, 0, "目次")
set_ph(s, 10, "本資料は Tableau Cloud から Snowflake 上の営業データを参照するための接続手順です")
body_lines(s, 11, [
    (0, "1. はじめに（目的・前提・全体構成）"),
    (0, "2. 役割分担"),
    (0, "3. 事前準備 － 貴社からご共有いただく情報"),
    (0, "4. 接続情報一覧"),
    (0, "5. Tableau Cloud での接続手順（①〜⑥）"),
    (0, "6. 利用上の注意"),
    (0, "7. トラブルシューティング"),
    (0, "8. PAT（アクセストークン）の運用"),
    (0, "付録. 弊社側 Snowflake 設定"),
])

# ============================================================================
# 3. はじめに
# ============================================================================
s = add(LY_OL)
set_ph(s, 0, "1. はじめに")
set_ph(s, 10, "Salesforce 由来の営業データ（Snowflake 上のマート）を Tableau Cloud から参照します")
body_lines(s, 11, [
    (0, "対象データ"),
    (1, "商談品目ワイドテーブル（1 行 = 商談 × 商品明細、全ディメンション結合済み）"),
    (0, "接続方式"),
    (1, "PAT（Programmatic Access Token）による認証。参照専用のサービスユーザー SVC_TABLEAU を使用"),
    (0, "前提"),
    (1, "Snowflake 環境（DB・ロール・マート）は構築済み"),
    (1, "Tableau は Tableau Cloud を利用"),
    (1, "Snowflake 側のユーザー・権限・トークン発行は弊社（エクスチュア）が実施"),
])

# ============================================================================
# 4. 全体構成
# ============================================================================
s = add(LY_STD)
set_ph(s, 0, "全体構成")
# 3ボックス + 矢印
def flowbox(l, title, sub):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(2.4), Inches(3.5), Inches(1.7))
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0xEE, 0xF3, 0xE9)
    b.line.color.rgb = GREEN; b.line.width = Pt(1)
    b.shadow.inherit = False
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = GRAY_D
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(10.5); r2.font.color.rgb = GRAY_D
def arrow(l):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(l), Inches(3.0), Inches(0.7), Inches(0.5))
    a.fill.solid(); a.fill.fore_color.rgb = GRAY_L; a.line.fill.background(); a.shadow.inherit = False
flowbox(0.7, "Salesforce", "営業データの源泉")
arrow(4.35)
flowbox(5.15, "Snowflake", "dbt でマートに整形\nSALES.ANALYTICS_MARTS")
arrow(8.8)
flowbox(9.6, "Tableau Cloud", "ダッシュボード作成・共有")
tb = textbox(s, 0.7, 4.6, 11.9, 1.8)
body_lines_tb = tb.text_frame
for i, (txt, b) in enumerate([
    ("Tableau Cloud → Snowflake の接続は サービスユーザー SVC_TABLEAU / PAT 認証 / 参照専用（SELECT のみ）", True),
    ("接続先: SALES.ANALYTICS_MARTS.GOLD_OPPORTUNITY_LINE_ITEM_WIDE（日本語カラム名）", False),
    ("Snowflake 側のアクセス元 IP はネットワークポリシーで許可リスト制限", False),
]):
    p = body_lines_tb.paragraphs[0] if i == 0 else body_lines_tb.add_paragraph()
    r = p.add_run(); r.text = ("• " + txt); r.font.size = Pt(12); r.font.color.rgb = GRAY_D; r.font.bold = b

# ============================================================================
# 5. 役割分担
# ============================================================================
s = add(LY_STD)
set_ph(s, 0, "2. 役割分担")
rows = [
    ["#", "作業", "担当"],
    ["1", "Tableau Cloud サイトの Pod（リージョン）と Egress IP レンジのご共有", "貴社"],
    ["2", "Snowflake 側のユーザー・権限・ネットワークポリシー作成", "弊社"],
    ["3", "PAT（トークン）の発行", "弊社"],
    ["4", "PAT・接続情報の受け渡し（安全な経路で）", "弊社 → 貴社"],
    ["5", "Tableau Cloud でのデータソース作成・パブリッシュ", "貴社"],
    ["6", "PAT 有効期限前のローテーションと再連携", "弊社"],
]
make_table(s, rows, 0.7, 1.3, 11.9, 3.6, col_w=[0.6, 8.3, 3.0])
tb = textbox(s, 0.7, 5.3, 11.9, 0.9)
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "本資料の主対象は #1・#5（貴社作業）です。#2〜#4 の内容は付録を参照ください。"
r.font.size = Pt(11.5); r.font.color.rgb = GRAY_D

# ============================================================================
# 6. 事前準備
# ============================================================================
s = add(LY_OL)
set_ph(s, 0, "3. 事前準備 － 貴社からご共有いただく情報")
set_ph(s, 10, "Snowflake 側で接続元 IP を許可リスト方式で制限するため、以下をご共有ください")
body_lines(s, 11, [
    (0, "① Tableau Cloud サイトの Pod 名 / デプロイ先リージョン"),
    (1, "サイト URL のサブドメイン（例: 10ax、us-east-1、eu-west-1a 等）"),
    (1, "管理者メニュー →「設定」→ サーバー情報 でも確認可能"),
    (0, "② 上記 Pod の Egress IP アドレス / CIDR レンジ"),
    (1, "Tableau が Pod 単位で公開（ヘルプ「IP アドレスとエンドポイントの許可リスト登録」）"),
    (0, "補足"),
    (1, "IP を絞らない運用（認証ポリシーで IP 要件を外す）も可能です。方針変更時はご相談ください"),
])

# ============================================================================
# 7. 接続情報一覧
# ============================================================================
s = add(LY_STD)
set_ph(s, 0, "4. 接続情報一覧")
rows = [
    ["項目", "値", "備考"],
    ["Snowflake アカウント URL", "弊社より別途ご連携", "例: xxxxx-xxxxx.snowflakecomputing.com"],
    ["ユーザー名", "SVC_TABLEAU", "固定"],
    ["パスワード欄に入れる値", "PAT トークン（弊社より別送）", "パスワードではなく PAT を入力"],
    ["ロール", "SALES_USER", ""],
    ["ウェアハウス", "SALES_WH", "自動起動 / 60 秒で自動停止"],
    ["データベース / スキーマ", "SALES / ANALYTICS_MARTS", ""],
    ["接続先テーブル", "GOLD_OPPORTUNITY_LINE_ITEM_WIDE", "日本語カラム名（例:「案件名」「合計金額」）"],
]
make_table(s, rows, 0.6, 1.25, 12.1, 4.0, col_w=[3.0, 4.4, 4.7], first_col_accent=True)
tb = textbox(s, 0.6, 5.5, 12.1, 0.8)
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "PAT は機密情報です。平文メールを避け、パスワード共有ツール等でお渡しします。有効期限は 90 日。"
r.font.size = Pt(11); r.font.color.rgb = RED

# ============================================================================
# 8-12. 手順
# ============================================================================
def step_slide(title, lead, items, shot_caption):
    s = add(LY_OL)
    set_ph(s, 0, title)
    set_ph(s, 10, lead)
    body_lines(s, 11, items)
    # 手順テキストの下にスクショ枠（中央）
    screenshot_ph(s, 3.15, 3.7, 7.0, 3.0, shot_caption)
    return s

step_slide(
    "5. 接続手順 ① 接続の開始",
    "Tableau Cloud（ブラウザ）または Tableau Desktop から作成できます",
    [
        (0, "Tableau Cloud にサインイン"),
        (0, "「新規」→「データソース」を選択"),
        (0, "コネクタ一覧から Snowflake を選択"),
    ],
    "（コネクタ選択画面）",
)
step_slide(
    "5. 接続手順 ② サーバーと認証",
    "認証は「ユーザー名とパスワード」を選び、パスワード欄に PAT を貼り付けます",
    [
        (0, "サーバー: 受領した Snowflake アカウント URL"),
        (0, "認証: ユーザー名とパスワード"),
        (1, "ユーザー名: SVC_TABLEAU"),
        (1, "パスワード: PAT トークンをそのまま貼り付け"),
        (0, "「サインイン」をクリック"),
        (1, "Snowflake の PAT はパスワードの代替。専用の認証方式は不要"),
    ],
    "（サーバー・認証情報の入力画面）",
)
step_slide(
    "5. 接続手順 ③ ロール・ウェアハウス",
    "接続後の詳細設定で指定します（プルダウンに表示されます）",
    [
        (0, "ロール: SALES_USER"),
        (0, "ウェアハウス: SALES_WH"),
    ],
    "（ロール / ウェアハウス選択）",
)
step_slide(
    "5. 接続手順 ④ テーブルの選択",
    "1 テーブルのみで完結します（結合不要）",
    [
        (0, "データベース: SALES"),
        (0, "スキーマ: ANALYTICS_MARTS"),
        (0, "テーブル: GOLD_OPPORTUNITY_LINE_ITEM_WIDE をキャンバスにドラッグ"),
    ],
    "（テーブル選択・キャンバス）",
)
s = add(LY_OL)
set_ph(s, 0, "5. 接続手順 ⑤ 接続モード ／ ⑥ パブリッシュ")
set_ph(s, 10, "PoC ではライブ接続を推奨します")
body_lines(s, 11, [
    (0, "⑤ 接続モード"),
    (1, "ライブ接続を推奨（データ量が少なく、SALES_WH は自動起動のため）"),
    (1, "抽出にする場合はパブリッシュ後にスケジュール更新を設定"),
    (0, "⑥ パブリッシュ"),
    (1, "データソース名を付けて Tableau Cloud にパブリッシュ"),
    (1, "認証情報は「埋め込みパスワード」（＝埋め込み PAT）を選択"),
    (1, "PAT ローテーション時は「接続の編集」から新トークンに更新"),
])

# ============================================================================
# 13. 利用上の注意
# ============================================================================
s = add(LY_OL)
set_ph(s, 0, "6. 利用上の注意")
set_ph(s, 10, "運用前に貴社ご担当者へ共有ください")
body_lines(s, 11, [
    (0, "カラム名が日本語です（「案件名」「クローズ日」「合計金額」など）"),
    (1, "GUI 操作は問題なし。カスタム SQL では列名を二重引用符で囲む"),
    (0, "SVC_TABLEAU は参照専用（SELECT のみ）。書き込み・DDL 不可"),
    (1, "参照範囲は SALES データベースのマート層のみ"),
    (0, "英語カラム名が必要な場合"),
    (1, "ビュー SALES.ANALYTICS_MARTS.V_OPPORTUNITY_LINE_ITEM（英語エイリアス）も利用可"),
])

# ============================================================================
# 14. トラブルシューティング
# ============================================================================
s = add(LY_STD)
set_ph(s, 0, "7. トラブルシューティング")
rows = [
    ["症状", "原因 / 対処"],
    ["Incorrect username or password", "PAT の期限切れ、または前後の空白・改行混入。弊社へ再発行を依頼"],
    ["ロールが選べない / not granted", "PAT のロール制限と不一致。弊社へ連絡"],
    ["IP address ... is not allowed to access", "送信元 IP が許可リスト未登録。エラー中の IP を弊社へ連携し追加を依頼"],
    ["ウェアハウスが起動しない / タイムアウト", "SALES_WH は自動起動。初回は数秒待ち。再試行で解消"],
    ["数値が想定と違う", "マート仕様差異の可能性。弊社と数値突合を実施"],
]
make_table(s, rows, 0.6, 1.3, 12.1, 3.8, col_w=[4.3, 7.8])

# ============================================================================
# 15. PAT の運用
# ============================================================================
s = add(LY_OL)
set_ph(s, 0, "8. PAT（アクセストークン）の運用")
set_ph(s, 10, "PAT には有効期限があり、定期的なローテーションが必要です")
body_lines(s, 11, [
    (0, "有効期限: 90 日"),
    (0, "ローテーション手順"),
    (1, "弊社が Snowflake で ROTATE を実行し、新トークンを発行"),
    (1, "旧トークンは 24 時間後に自動失効"),
    (1, "新トークンを貴社へ再連携"),
    (1, "貴社が Tableau データソースの「接続の編集」で新トークンに更新"),
    (0, "期限切れの兆候: データ更新失敗 / サインインエラー → 弊社へご連絡ください"),
])

# ============================================================================
# 16. 付録：弊社側 Snowflake 設定
# ============================================================================
s = add(LY_STD)
set_ph(s, 0, "付録. 弊社側 Snowflake 設定（Snowflake/sql/06_users.sql）")
code = (
    "USE ROLE SECURITYADMIN;\n"
    "\n"
    "-- 1) サービスユーザー\n"
    "CREATE USER IF NOT EXISTS SVC_TABLEAU\n"
    "  TYPE = SERVICE  DEFAULT_ROLE = SALES_USER  DEFAULT_WAREHOUSE = SALES_WH;\n"
    "GRANT ROLE SALES_USER TO USER SVC_TABLEAU;   -- 参照権限は SALES_USER → SALES__R\n"
    "\n"
    "-- 2) ネットワークポリシー（PAT の前提。貴社共有の Egress IP を許可）\n"
    "CREATE NETWORK POLICY IF NOT EXISTS TABLEAU_PAT_NP\n"
    "  ALLOWED_IP_LIST = ('<貴社共有の IP/CIDR>');\n"
    "ALTER USER SVC_TABLEAU SET NETWORK_POLICY = TABLEAU_PAT_NP;\n"
    "\n"
    "-- 3) PAT 発行（Snowsight で手動実行。token_secret は発行時のみ表示）\n"
    "ALTER USER IF EXISTS SVC_TABLEAU ADD PROGRAMMATIC ACCESS TOKEN TABLEAU_PAT\n"
    "  ROLE_RESTRICTION = 'SALES_USER'  DAYS_TO_EXPIRY = 90;\n"
    "\n"
    "-- ローテーション\n"
    "ALTER USER IF EXISTS SVC_TABLEAU ROTATE PROGRAMMATIC ACCESS TOKEN TABLEAU_PAT\n"
    "  EXPIRE_ROTATED_TOKEN_AFTER_HOURS = 24;"
)
tb = textbox(s, 0.6, 1.2, 12.1, 5.4)
tb.fill.solid(); tb.fill.fore_color.rgb = RGBColor(0xF6, 0xF6, 0xF6)
tb.line.color.rgb = GRAY_L; tb.line.width = Pt(0.75)
tf = tb.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.15); tf.margin_top = Inches(0.12)
for i, line in enumerate(code.split("\n")):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = line if line else " "
    r.font.name = "Consolas"; r.font.size = Pt(10); r.font.color.rgb = GRAY_D

# ============================================================================
# 17. 最終ページ
# ============================================================================
add(LY_END)

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides))
